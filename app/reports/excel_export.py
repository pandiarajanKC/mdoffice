"""Generic Excel export shared by every report (master spec section 39:
"Allow export to Excel where practical").
"""
from __future__ import annotations

import io

from flask import Response
from openpyxl import Workbook
from openpyxl.styles import Font

_HEADER_FONT = Font(bold=True)


def export_to_excel(*, title: str, columns: list[dict], rows: list[dict]) -> Response:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] or "Report"  # Excel sheet name limit

    for col_idx, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col["label"])
        cell.font = _HEADER_FONT

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, col in enumerate(columns, start=1):
            value = row.get(col["key"])
            if col["type"] in ("date",) and value is not None:
                value = value.strftime("%d-%b-%Y")
            elif value is None:
                value = ""
            ws.cell(row=row_idx, column=col_idx, value=value)

    for col_idx, col in enumerate(columns, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max(14, len(col["label"]) + 2)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"{title.replace(' ', '_')}.xlsx"
    return Response(
        buffer.read(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
